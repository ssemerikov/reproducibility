#!/usr/bin/env python3
"""Audit the supplied curriculum workbook; emit aggregate comparisons only.

This is an inclusion assessment, not a replacement for the thesis extraction.
Neither respondent rows nor identifying/free-text fields are exported.
"""
import argparse
import csv
import hashlib
import json
from collections import Counter
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
REPORT = ROOT / 'revision/reports'
LABELS = {'зовсім не важливо': -1, 'мало важливо': 1,
          'утруднююсь відповісти': 0, 'скоріше, важливо': 2,
          'скоріше важливо': 2, 'дуже важливо': 3}
WEIGHTS = {'маю початкові уявлення': 1,
           'обізнаний з окремими складовими': 2,
           'глибоко обізнаний з окремими складовими': 3, 'експерт': 4}
COUNT_KEYS = ['not_important', 'slightly_important', 'unable_to_answer',
              'rather_important', 'very_important']
LEVELS = [-1, 1, 0, 2, 3]


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--workbook', type=Path, default=ROOT / 'Квантову інформатику - в школу! (Responses).xlsx')
    parser.add_argument('--data', type=Path, default=ROOT / 'revision/data')
    parser.add_argument('--output', type=Path, default=ROOT / 'revision/data')
    args = parser.parse_args()
    path = args.workbook
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    formulas = openpyxl.load_workbook(path, read_only=True, data_only=False)
    sheet = list(book['Form Responses 1'].values)
    header = sheet[0]
    rows = [(i, r) for i, r in enumerate(sheet[1:], 2) if any(v is not None for v in r)]
    digits = list(book['digits'].values)
    coded_formulas = list(formulas['digits'].values)
    columns = [i for i, v in enumerate(header) if str(v).startswith('Оцінить важливість')]
    with (args.data / 'curriculum_ratings.csv').open() as stream:
        thesis = list(csv.DictReader(stream))
    assert len(rows) == 36 and len(columns) == len(thesis) == 64
    assert all(header[c] == digits[0][c] for c in columns)
    assert all(WEIGHTS[r[81]] == digits[i - 1][81] for i, r in rows)
    items = []
    stratified = []
    mismatches = 0
    for column, original in zip(columns, thesis):
        counts = Counter(LABELS[r[column]] for _, r in rows)
        stored_counts = Counter(digits[i - 1][column] for i, _ in rows)
        score = sum(LABELS[r[column]] * WEIGHTS[r[81]] for _, r in rows)
        stored_score = sum(digits[i - 1][column] * digits[i - 1][81] for i, _ in rows)
        assert [stored_counts[x] for x in LEVELS] == [int(original[k]) for k in COUNT_KEYS]
        assert stored_score == int(original['reported_expertise_weighted_score'])
        different = 0
        for i, row in rows:
            coordinate = f'{openpyxl.utils.get_column_letter(column + 1)}{i}'
            formula = coded_formulas[i - 1][column]
            assert f"'Form Responses 1'!{coordinate}=" in formula
            if LABELS[row[column]] != digits[i - 1][column]:
                # Verify the proposed cause for every observed mismatch.
                assert row[column] == 'скоріше важливо'
                assert digits[i - 1][column] == 3
                assert "'Опрацювання'!$C$5" in formula
                assert formula.endswith("'Опрацювання'!$E$6)))))")
                different += 1
        mismatches += different
        item = {'item_id': original['item_id'], 'item_uk': original['item_uk'],
                'source_column': openpyxl.utils.get_column_letter(column + 1),
                'n': len(rows), 'recoded_weighted_score': score,
                'thesis_weighted_score': stored_score,
                'recoded_response_count': different}
        for key, value in zip(COUNT_KEYS, LEVELS):
            item['raw_' + key] = counts[value]
            item['thesis_' + key] = int(original[key])
        items.append(item)
        for weight in sorted(set(WEIGHTS[r[81]] for _, r in rows)):
            for level in LEVELS:
                stratified.append({'item_id': original['item_id'], 'expertise_weight': weight,
                    'rating_code': level,
                    'raw_count': sum(WEIGHTS[r[81]] == weight and LABELS[r[column]] == level for _, r in rows),
                    'stored_count': sum(WEIGHTS[r[81]] == weight and digits[i - 1][column] == level for i, r in rows)})
    assert mismatches == 63
    changed = [r for r in items if r['recoded_response_count']]
    assert [r['item_id'] for r in changed] == ['2.01', '2.02', '2.03', '2.04']
    scores = [r['recoded_weighted_score'] for r in items]
    cutoff = min(scores) + 0.45 * (max(scores) - min(scores))
    for row in items:
        row['thesis_cutoff_met'] = row['thesis_weighted_score'] >= 120.8
        row['recoded_fixed_cutoff_met'] = row['recoded_weighted_score'] >= 120.8
        row['recoded_recomputed_cutoff_met'] = row['recoded_weighted_score'] >= cutoff
    summary = {
        'assessment_date': '2026-10-13', 'file': path.name, 'sha256': digest,
        'sheets': book.sheetnames, 'respondents': len(rows), 'rating_items': len(columns),
        'response_date_range': [min(r[0] for _, r in rows).date().isoformat(),
                                max(r[0] for _, r in rows).date().isoformat()],
        'missing_item_ratings': sum(r[c] is None for _, r in rows for c in columns),
        'domain_importance_items': 8,
        'missing_domain_ratings': sum(r[c] is None for _, r in rows for c in range(73, 81)),
        'self_rated_expertise_counts_by_weight': dict(Counter(WEIGHTS[r[81]] for _, r in rows)),
        'numeric_sheet_matches_thesis_counts_and_scores_items': 64,
        'raw_recoding_matches_thesis_items': 60,
        'raw_numeric_response_mismatches': mismatches,
        'affected_items': changed,
        'coding_cause': 'The no-comma label скоріше важливо fails the exact comparison with скоріше, важливо and falls through to code 3 instead of 2.',
        'recoding_assumption': 'Treat the two punctuation variants as the same response category (rather important). Original workbook unchanged.',
        'selection': {'published_cutoff': 120.8, 'published_count': 30,
                      'recoded_at_published_cutoff_count': sum(s >= 120.8 for s in scores),
                      'recoded_min': min(scores), 'recoded_max': max(scores),
                      'recoded_range_rule_cutoff': cutoff,
                      'recoded_at_recomputed_cutoff_count': sum(s >= cutoff for s in scores),
                      'published_shaded_selection_count': 41},
        'source_confirmation': 'On 13 October 2026 the supplying author confirmed this as the original curriculum-survey export used for the thesis. Ethics approval and informed consent are reported in the article's declarations.',
        'analysis_scope': 'Audit of the existing 36-person curriculum survey; no independent sample and no silent replacement of published aggregates.',
        'limits': ['No student outcomes or separate 26-teacher survey in this workbook.',
                   'Does not resolve the two student-entry chart/table conflicts.',
                   'Presence of contact, job, tenure and free-text fields precludes distributing the original as an anonymous aggregate supplement.',
                   'Provenance is author-confirmed and numerically corroborated; numerical agreement does not establish ethics/consent documentation.',
                   'Self-rated expertise weights are not independently measured expertise.'],
    }
    args.output.mkdir(parents=True, exist_ok=True)
    (args.output / 'workbook_assessment.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2) + '\n')
    with (args.output / 'workbook_item_comparison.csv').open('w', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=list(items[0]))
        writer.writeheader()
        writer.writerows(items)
    with (args.output / 'curriculum_weighted_counts.csv').open('w', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=list(stratified[0]))
        writer.writeheader()
        writer.writerows(stratified)
    assert hashlib.sha256(path.read_bytes()).hexdigest() == digest
    print(json.dumps({k: summary[k] for k in ['respondents', 'rating_items',
          'raw_recoding_matches_thesis_items', 'raw_numeric_response_mismatches', 'selection']}, indent=2))


if __name__ == '__main__':
    main()
